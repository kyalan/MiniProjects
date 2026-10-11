import io
from datetime import date, datetime, timedelta, timezone

from django.test import SimpleTestCase
from openpyxl import Workbook

from venues.equipment import equipment_category
from venues.pipeline.coords import dms_to_decimal, locate, manual_location, parse_sizes
from venues.pipeline.listing import html_text, parse_listing, parse_room_page
from venues.pipeline.timetable import parse_xlsx
from venues.present import classify_window


class ParserTests(SimpleTestCase):
    def test_coordinates(self):
        self.assertAlmostEqual(dms_to_decimal("114-9-33"), 114 + 9 / 60 + 33 / 3600)
        self.assertAlmostEqual(dms_to_decimal("114-9-60"), 114 + 10 / 60)
        lat, lng, quality = locate("22-16-38", "114-9-33")
        self.assertEqual(quality, "ok")
        self.assertIsNotNone(lat)
        self.assertIsNotNone(lng)
        self.assertEqual(locate("00-00-00", "00-00-00")[2], "missing")
        self.assertEqual(locate("22-33-04", "144-14-90")[2], "invalid")
        self.assertEqual(locate("", "")[2], "missing")
        self.assertEqual(manual_location("Hiu Kwong Street Sports Centre"), (22.3197111, 114.2251903))
        self.assertEqual(manual_location("Sham Shui Po Sports Centre"), (22.3308265, 114.146248))
        self.assertEqual(manual_location("Tung Cheong Street Sports Centre"), (22.4496581, 114.1678594))
        self.assertEqual(manual_location("Choi Wing Road Sports Centre"), (22.3294345, 114.213932))
        self.assertEqual(manual_location("Ma Chai Hang Sports Centre"), (22.3445776, 114.1867359))
        self.assertIsNone(manual_location("Aberdeen Sports Centre"))

    def test_sizes(self):
        self.assertEqual(parse_sizes("75m 2 50m 2 "), [75.0, 50.0])
        self.assertEqual(parse_sizes("317 m2"), [317.0])
        self.assertEqual(parse_sizes("121m 2 "), [121.0])

    def test_listing_regions_and_access(self):
        html = """
        <table class="fitness_tb">
          <thead><tr><th><a name="dist1"></a>Hong Kong Island</th></tr></thead>
          <tbody>
            <tr>
              <td class="venu">Aberdeen Sports Centre</td>
              <td class="addr">Aberdeen</td>
              <td class="tel">2555 8909</td>
              <td class="detail"><a href="/clpss/en/webApp/FitnessRoomDetails.do?id=1374">Details</a></td>
            </tr>
            <tr>
              <td class="venu">Chai Wan Sports Centre</td>
              <td class="addr">Chai Wan</td>
              <td class="tel">2897 9144</td>
              <td class="detail"><a href="/clpss/en/webApp/FitnessRoomDetails.do?id=1365">Details</a>*</td>
            </tr>
          </tbody>
        </table>
        """
        rows = parse_listing(html)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["detail_id"], 1374)
        self.assertEqual(rows[0]["address"], "Aberdeen")
        self.assertEqual(rows[0]["phone"], "2555 8909")
        self.assertEqual(rows[0]["region"], "hk_island")
        self.assertFalse(rows[0]["accessible"])
        self.assertTrue(rows[1]["accessible"])
        page = """
        <h2>Size of Fitness Room : 279m<sup>2</sup></h2>
        <table><tr><th>Equipment</th><th>No. (Set)</th></tr>
        <tr><td><span>Treadmill </span></td><td><span>3</span></td></tr>
        <tr><td>Arc Trainer</td><td>1</td></tr></table>
        """
        parsed_page = parse_room_page(page)
        self.assertIn("279", parsed_page["size"])
        self.assertEqual(parsed_page["equipment"][0]["name"], "Treadmill")
        self.assertEqual(parsed_page["equipment"][0]["quantity"], 3)
        self.assertEqual(parsed_page["equipment"][1]["name"], "Arc Trainer")

    def test_xlsx_timetable(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet["A9"] = "時票名額 Quota for Hourly Ticket Users: 6"
        sheet["A10"] = "月票名額 Quota for Monthly Ticket Users: 22"
        sheet["A12"] = "日期 Date\n時間 Time"
        for day in range(1, 32):
            sheet.cell(12, day + 1, f"{day}\n週一\nMon")
        for row, label in enumerate(("07:00 - 08:00", "22:00 - 23:00"), start=13):
            sheet.cell(row, 1, label)
            for day in range(1, 32):
                sheet.cell(row, day + 1, "A" if day % 2 else "M")
        sheet["A31"] = "發出日期 Date of issue 21.9.2026"
        sheet["A32"] = "最新更新日期 Date of latest update 20.9.2026"
        buffer = io.BytesIO()
        workbook.save(buffer)
        parsed = parse_xlsx(buffer.getvalue(), 2026, 10)
        self.assertEqual(len(parsed), 1)
        grid = parsed[0]
        self.assertEqual(grid["hourly_quota"], 6)
        self.assertEqual(grid["monthly_quota"], 22)
        self.assertEqual(grid["issued_on"], date(2026, 9, 21))
        self.assertEqual(grid["source_updated_on"], date(2026, 9, 20))
        self.assertEqual(len(grid["slots"]), 62)
        self.assertEqual(grid["slots"][0]["status"], "A")
        self.assertEqual(grid["slots"][0]["date"], date(2026, 10, 1))
        self.assertEqual(grid["slots"][1]["status"], "M")

    def test_excel_date_serials_and_html_sizes(self):
        from datetime import datetime

        workbook = Workbook()
        sheet = workbook.active
        sheet["A12"] = "日期 Date"
        for day in range(1, 8):
            sheet.cell(12, day + 1, datetime(1900, 1, day))
        sheet["A13"] = "07:00 - 08:00"
        for day in range(1, 8):
            sheet.cell(13, day + 1, "A")
        buffer = io.BytesIO()
        workbook.save(buffer)
        parsed = parse_xlsx(buffer.getvalue(), 2026, 10)
        self.assertEqual(len(parsed[0]["slots"]), 7)
        self.assertEqual(parsed[0]["slots"][0]["date"], date(2026, 10, 1))
        self.assertEqual(html_text("165m<sup>2</sup>"), "165m2")
        self.assertEqual(html_text("75m<sup>2</sup><br>50m<sup>2</sup>"), "75m2 50m2")

    def test_two_hour_window_and_equipment_groups(self):
        start = datetime(2026, 10, 11, 18, 0, tzinfo=timezone.utc)
        end = start + timedelta(hours=2)
        hour = timedelta(hours=1)
        self.assertEqual(classify_window([], start, end), "other")
        available = [(start, start + hour, "A"), (start + hour, end, "A")]
        self.assertEqual(classify_window(available, start, end), "open")
        self.assertEqual(classify_window([(start, start + hour, "A")], start, end), "soon")
        occupied = [(start, start + hour, "A"), (start + hour, end, "P")]
        self.assertEqual(classify_window(occupied, start, end), "soon")
        self.assertEqual(equipment_category("Electronic Treadmill"), "treadmill")
        self.assertEqual(equipment_category("Recumbent bike"), "bike")
        self.assertEqual(equipment_category("Elliptical Cross Trainer"), "elliptical")
        self.assertEqual(equipment_category("Natural Runner Plus"), "elliptical")
        self.assertEqual(equipment_category("Natural Runner"), "elliptical")
        self.assertEqual(equipment_category("Arc Trainer"), "elliptical")
        self.assertEqual(equipment_category("ARC Trainer"), "elliptical")
        self.assertEqual(equipment_category("Arc Trainer Total Body"), "elliptical")
        self.assertEqual(equipment_category("Stretch Trainer"), "other")
        self.assertEqual(equipment_category("Rowing Machine"), "rowing")
        self.assertEqual(equipment_category("Flat / Incline Bench"), "free_weights")
        self.assertEqual(equipment_category("Plate Loaded Smith Press"), "free_weights")
        self.assertEqual(equipment_category("Chest Press"), "upper_body")
        self.assertEqual(equipment_category("Leg Press"), "lower_body")
        self.assertEqual(equipment_category("Leg Raise"), "core")
        self.assertEqual(equipment_category("Dual Adjustable Pulley"), "cable")
        self.assertEqual(equipment_category("Stretching Mat"), "other")
