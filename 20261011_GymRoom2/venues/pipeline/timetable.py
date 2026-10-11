import io
import re
from datetime import date, datetime, time

from openpyxl import load_workbook

STATUSES = {"A", "P", "T", "B", "M", "S"}
TIME_RE = re.compile(r"(\d{1,2}:\d{2})\s*[-–]\s*(\d{1,2}:\d{2})")
DAY_RE = re.compile(r"(\d{1,2})(?:\s|$|\n|週|周)")


def parse_clock(value: str) -> time:
    parsed = datetime.strptime(value, "%H:%M")
    return time(parsed.hour, parsed.minute)


def day_of(value) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, datetime):
        # Some LCSD sheets store the day number as an Excel date serial.
        if (value.year == 1900 and value.month == 1) or value.year >= 2000:
            return value.day if 1 <= value.day <= 31 else None
        return None
    if isinstance(value, (int, float)):
        day = int(value)
        return day if 1 <= day <= 31 else None
    text = str(value).strip()
    if not text or TIME_RE.search(text):
        return None
    match = DAY_RE.match(text)
    if not match:
        return None
    day = int(match.group(1))
    return day if 1 <= day <= 31 else None


def normalize_status(value) -> str:
    if value is None:
        return ""
    text = str(value).strip().upper()
    return text if text in STATUSES else ""


def _dmy(blob: str, pattern: str) -> date | None:
    match = re.search(pattern, blob, flags=re.IGNORECASE)
    if not match:
        return None
    day, month, year = (int(part) for part in match.groups())
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _quota(blob: str, english: str, chinese: str) -> int | None:
    match = re.search(rf"{english}\s*:\s*(\d+)", blob, flags=re.IGNORECASE)
    if not match:
        match = re.search(rf"{chinese}\s*[:：]\s*(\d+)", blob)
    if not match:
        return None
    return int(match.group(1))


def parse_grid(rows: list[list], year: int, month: int, room_no: int = 1) -> dict | None:
    texts: list[str] = []
    best = 0
    day_cols: dict[int, int] = {}
    header_idx = None
    for index, row in enumerate(rows):
        found: dict[int, int] = {}
        for col, value in enumerate(row):
            if value is not None and not isinstance(value, (int, float)):
                texts.append(str(value))
            day = day_of(value)
            if day is not None:
                found[col] = day
        if len(found) > best:
            best = len(found)
            day_cols = found
            header_idx = index
    if best < 7 or header_idx is None:
        return None
    unique_days: dict[int, int] = {}
    seen_days: set[int] = set()
    for col, day in day_cols.items():
        if day in seen_days:
            continue
        seen_days.add(day)
        unique_days[col] = day
    day_cols = unique_days

    blob = "\n".join(texts)
    slots = []
    for row in rows[header_idx + 1 :]:
        found_time = None
        for value in row:
            if value is None:
                continue
            found_time = TIME_RE.search(str(value))
            if found_time:
                break
        if not found_time:
            continue
        start = parse_clock(found_time.group(1))
        end = parse_clock(found_time.group(2))
        for col, day in day_cols.items():
            if col >= len(row):
                continue
            status = normalize_status(row[col])
            if not status:
                continue
            try:
                slot_date = date(year, month, day)
            except ValueError:
                continue
            slots.append(
                {
                    "date": slot_date,
                    "start": start,
                    "end": end,
                    "status": status,
                }
            )
    if not slots:
        return None
    return {
        "room_no": room_no,
        "hourly_quota": _quota(blob, r"Hourly Ticket Users", "時票名額"),
        "monthly_quota": _quota(blob, r"Monthly Ticket Users", "月票名額"),
        "issued_on": _dmy(blob, r"Date of issue\s+(\d{1,2})\.(\d{1,2})\.(\d{4})"),
        "source_updated_on": _dmy(blob, r"Date of latest update\s+(\d{1,2})\.(\d{1,2})\.(\d{4})"),
        "slots": slots,
    }


def parse_xlsx(content: bytes, year: int, month: int) -> list[dict]:
    if not content.startswith(b"PK"):
        raise ValueError("not an xlsx")
    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    parsed: list[dict] = []
    try:
        for index, worksheet in enumerate(workbook.worksheets, start=1):
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            grid = parse_grid(rows, year, month, room_no=index)
            if grid:
                parsed.append(grid)
    finally:
        workbook.close()
    return parsed


def parse_pdf(content: bytes, year: int, month: int) -> list[dict]:
    import pdfplumber

    found: list[dict] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables() or []:
                rows = [list(row or []) for row in table]
                grid = parse_grid(rows, year, month, room_no=1)
                if grid:
                    found.append(grid)
            if found:
                continue
            fallback = _parse_pdf_text(page.extract_text() or "", year, month)
            if fallback:
                found.append(fallback)
    return found[:1]


def _parse_pdf_text(text: str, year: int, month: int) -> dict | None:
    slots = []
    for line in text.splitlines():
        match = TIME_RE.search(line)
        if not match:
            continue
        rest = line[match.end() :]
        letters = [token for token in rest.split() if token in STATUSES]
        if len(letters) < 7:
            continue
        start = parse_clock(match.group(1))
        end = parse_clock(match.group(2))
        for offset, status in enumerate(letters, start=1):
            try:
                slot_date = date(year, month, offset)
            except ValueError:
                break
            slots.append({"date": slot_date, "start": start, "end": end, "status": status})
    if not slots:
        return None
    return {
        "room_no": 1,
        "hourly_quota": _quota(text, r"Hourly Ticket Users", "時票名額"),
        "monthly_quota": _quota(text, r"Monthly Ticket Users", "月票名額"),
        "issued_on": _dmy(text, r"Date of issue\s+(\d{1,2})\.(\d{1,2})\.(\d{4})"),
        "source_updated_on": _dmy(text, r"Date of latest update\s+(\d{1,2})\.(\d{1,2})\.(\d{4})"),
        "slots": slots,
    }
